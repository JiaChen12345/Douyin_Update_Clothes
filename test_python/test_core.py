import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from douyin_automation_dual_store_v49 import DouyinDraftCreator, ExcelStore, TaobaoAdapter, append_urls, collect_images, collect_videos, extension_manifest_dir, extract_zip_safe, load_queue_metadata, natural_key


class CoreTests(unittest.TestCase):
    def test_extension_manifest_dir_accepts_root_and_version_parent(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            version = root / "1.2.3_0"
            version.mkdir()
            (version / "manifest.json").write_text("{}", encoding="utf-8")
            self.assertEqual(extension_manifest_dir(version), version.resolve())
            self.assertEqual(extension_manifest_dir(root), version.resolve())

    def test_natural_sort(self):
        self.assertEqual(sorted(["10.jpg", "2.jpg", "1.jpg"], key=natural_key), ["1.jpg", "2.jpg", "10.jpg"])

    def test_hover_icon_must_be_inside_image_card(self):
        image = {"x": 100, "y": 50, "width": 140, "height": 170}
        trash = {"x": 210, "y": 185, "width": 24, "height": 24}
        outside = {"x": 260, "y": 185, "width": 24, "height": 24}
        self.assertTrue(DouyinDraftCreator._box_inside(trash, image))
        self.assertFalse(DouyinDraftCreator._box_inside(outside, image))

    def test_collect_images_and_safe_extract(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "素材.zip"
            with zipfile.ZipFile(archive, "w") as output:
                output.writestr("商品主图/2.jpg", b"x")
                output.writestr("商品主图/10.jpg", b"x")
                output.writestr("商品详情图/1.png", b"x")
            extracted = extract_zip_safe(archive, root / "out")
            self.assertEqual([p.name for p in collect_images(extracted, "主图")], ["2.jpg", "10.jpg"])
            self.assertEqual([p.name for p in collect_images(extracted, "详情图")], ["1.png"])

    def test_collect_taobao_main_images_by_filename(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "商品文件夹").mkdir()
            (root / "商品文件夹" / "主图01.jpg").write_bytes(b"x")
            (root / "商品文件夹" / "普通图.jpg").write_bytes(b"x")
            self.assertEqual([p.name for p in collect_images(root, "主图")], ["主图01.jpg"])

    def test_taobao_adapter_domains(self):
        adapter = TaobaoAdapter({})
        self.assertTrue(adapter.matches("https://item.taobao.com/item.htm?id=1"))
        self.assertTrue(adapter.matches("https://detail.tmall.com/item.htm?id=1"))
        self.assertFalse(adapter.matches("https://example.com/item/1"))

    def test_collect_videos_naturally(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "video10.mp4").write_bytes(b"x")
            (root / "video2.mov").write_bytes(b"x")
            (root / "ignore.txt").write_bytes(b"x")
            self.assertEqual([p.name for p in collect_videos(root)], ["video2.mov", "video10.mp4"])

    def test_excel_status(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "tasks.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["网址"])
            sheet.append(["https://cs.17zwd.com/item/1"])
            book.save(path)
            store = ExcelStore(path)
            store.load()
            store.update(2, {"站点": "17网", "状态": "草稿成功"})
            reopened = ExcelStore(path)
            reopened.load()
            self.assertEqual(reopened.tasks()[0]["status"], "草稿成功")

    def test_single_store_completion_deletes_after_store_one(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "tasks.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["网址", "抖店1状态", "抖店2状态"])
            sheet.append(["https://cs.17zwd.com/item/1", "草稿成功", ""])
            book.save(path)
            store = ExcelStore(path)
            store.load()
            self.assertEqual(store.delete_completed_rows("single"), 1)

    def test_dual_store_completion_requires_both(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "tasks.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["网址", "抖店1状态", "抖店2状态"])
            sheet.append(["https://cs.17zwd.com/item/1", "草稿成功", ""])
            book.save(path)
            store = ExcelStore(path)
            store.load()
            self.assertEqual(store.delete_completed_rows("dual"), 0)

    def test_queue_persists_category_and_destination(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "task_queue.json"
            added, invalid = append_urls(path, ["https://item.taobao.com/item.htm?id=1"], "knitwear", "zwd17")
            self.assertEqual((added, invalid), (1, []))
            task = load_queue_metadata(path)["https://item.taobao.com/item.htm?id=1"]
            self.assertEqual(task["category"], "knitwear")
            self.assertEqual(task["destination"], "zwd17")

    def test_old_queue_record_gets_safe_defaults(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "task_queue.json"
            path.write_text('[{"url":"https://item.taobao.com/item.htm?id=2"}]', encoding="utf-8")
            task = load_queue_metadata(path)["https://item.taobao.com/item.htm?id=2"]
            self.assertEqual(task["category"], "tshirt")
            self.assertEqual(task["destination"], "douyin")


if __name__ == "__main__":
    unittest.main()
